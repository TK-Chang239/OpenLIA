"""Phase 9 — attachment persistence service.

Validates uploads (size, count, mime), writes bytes via the storage
backend, extracts text for non-native-content formats, and persists
``ChatAttachment`` rows linked to a parent ``ChatMessage``.

These tests exercise the public surface (``validate_uploads``,
``persist_attachments``) and assert on observable outcomes — DB rows,
files on disk, extracted text content. Internal mime-table layout and
parser choice are not asserted.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openlia_server.db.models.content import ChatMessage, ChatSession
from openlia_server.services import attachment_storage  # noqa: F401  (autouse env)
from openlia_server.services.attachments import (
    FileUpload,
    persist_attachments,
    validate_uploads,
)
from sqlalchemy.orm import Session


@pytest.fixture(autouse=True)
def _isolated_storage(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setenv("OPENLIA_ATTACHMENTS_DIR", str(tmp_path / "attachments"))
    return tmp_path


@pytest.fixture
def db(db_session: Session) -> Session:
    """Reuse the project-wide ``db_session`` fixture and seed the
    parent ChatSession + ChatMessage rows that ``message_id="m1"`` points at.
    The ``conftest.py`` ``_seed_test_users`` autouse fixture provides ``u-1``
    so we attach the session to that user.
    """
    session_row = ChatSession(
        id="s1",
        user_id="u-1",
        department="secretary",
        title="t",
    )
    msg = ChatMessage(
        id="m1",
        session_id="s1",
        role="user",
        content="hi",
        created_at=datetime.now(UTC),
    )
    db_session.add_all([session_row, msg])
    db_session.commit()
    return db_session


def _txt_upload(name: str = "notes.txt", body: bytes = b"hello world") -> FileUpload:
    return FileUpload(filename=name, mime_type="text/plain", content=body)


# ─── validate_uploads ────────────────────────────────────────────────────────


def test_validate_passes_supported_text_file() -> None:
    assert validate_uploads([_txt_upload()]) == []


def test_validate_rejects_oversized_file() -> None:
    big = FileUpload(filename="big.txt", mime_type="text/plain", content=b"x" * (26 * 1024 * 1024))
    errors = validate_uploads([big])
    assert len(errors) == 1
    assert errors[0].filename == "big.txt"
    assert errors[0].reason == "file_too_large"


def test_validate_rejects_too_many_files() -> None:
    files = [_txt_upload(f"f{i}.txt") for i in range(11)]
    errors = validate_uploads(files)
    assert any(e.reason == "too_many_files" for e in errors)


def test_validate_rejects_disallowed_mime() -> None:
    upload = FileUpload(filename="evil.zip", mime_type="application/zip", content=b"PK\x03\x04...")
    errors = validate_uploads([upload])
    assert errors[0].reason == "type_not_allowed"


def test_validate_returns_per_file_errors() -> None:
    """Multiple problematic files each get their own error."""
    files = [
        _txt_upload("ok.txt"),
        FileUpload(filename="bad.zip", mime_type="application/zip", content=b"PK"),
        FileUpload(
            filename="huge.txt",
            mime_type="text/plain",
            content=b"x" * (26 * 1024 * 1024),
        ),
    ]
    errors = validate_uploads(files)
    by_filename = {e.filename: e.reason for e in errors}
    assert "ok.txt" not in by_filename
    assert by_filename["bad.zip"] == "type_not_allowed"
    assert by_filename["huge.txt"] == "file_too_large"


# ─── persist_attachments ─────────────────────────────────────────────────────


def test_persist_text_file_creates_row_with_extracted_text(db: Session) -> None:
    rows = persist_attachments(db, message_id="m1", uploads=[_txt_upload(body=b"hello world")])
    assert len(rows) == 1
    row = rows[0]
    assert row.message_id == "m1"
    assert row.filename == "notes.txt"
    assert row.mime_type == "text/plain"
    assert row.extracted_text == "hello world"
    assert row.extracted_at is not None
    assert Path(row.storage_path).is_file()
    assert Path(row.storage_path).read_bytes() == b"hello world"


def test_persist_image_does_not_extract_text(db: Session) -> None:
    upload = FileUpload(filename="x.png", mime_type="image/png", content=b"\x89PNG fake")
    [row] = persist_attachments(db, message_id="m1", uploads=[upload])
    assert row.extracted_text is None
    assert row.extracted_at is None


def test_persist_pdf_extracts_text(db: Session) -> None:
    """Generate a tiny real PDF on the fly and verify extraction returns
    something resembling the embedded text."""
    pdf_bytes = _make_minimal_pdf("hello pdf body")
    upload = FileUpload(filename="r.pdf", mime_type="application/pdf", content=pdf_bytes)
    [row] = persist_attachments(db, message_id="m1", uploads=[upload])
    assert row.extracted_text is not None
    assert "hello pdf body" in row.extracted_text


def test_persist_docx_extracts_text(db: Session) -> None:
    docx_bytes = _make_minimal_docx("docx body line 1")
    upload = FileUpload(
        filename="m.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=docx_bytes,
    )
    [row] = persist_attachments(db, message_id="m1", uploads=[upload])
    assert row.extracted_text is not None
    assert "docx body line 1" in row.extracted_text


def test_persist_xlsx_extracts_text(db: Session) -> None:
    xlsx_bytes = _make_minimal_xlsx({"Sheet1": [["A1", "B1"], ["A2", "B2"]]})
    upload = FileUpload(
        filename="model.xlsx",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=xlsx_bytes,
    )
    [row] = persist_attachments(db, message_id="m1", uploads=[upload])
    assert row.extracted_text is not None
    assert "A1" in row.extracted_text
    assert "B2" in row.extracted_text


def test_persist_writes_distinct_storage_paths_for_duplicate_filenames(
    db: Session,
) -> None:
    rows = persist_attachments(
        db,
        message_id="m1",
        uploads=[_txt_upload("dup.txt", b"a"), _txt_upload("dup.txt", b"b")],
    )
    assert rows[0].storage_path != rows[1].storage_path
    assert Path(rows[0].storage_path).read_bytes() == b"a"
    assert Path(rows[1].storage_path).read_bytes() == b"b"


def test_persist_records_size_bytes_from_content(db: Session) -> None:
    upload = _txt_upload(body=b"abcde")
    [row] = persist_attachments(db, message_id="m1", uploads=[upload])
    assert row.size_bytes == 5


# ─── helpers: minimal real-format file generators ───────────────────────────


def _make_minimal_pdf(text: str) -> bytes:
    """Build a single-page PDF with ``text`` as content. Uses pypdf's writer."""
    from pypdf import PdfWriter
    from pypdf.generic import (
        DecodedStreamObject,
        DictionaryObject,
        NameObject,
    )

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    page = writer.pages[0]

    # Inject a content stream that draws the text. pypdf's content_stream API
    # is finicky; the simplest portable path is to set a stream with raw PDF
    # operators.
    content = DecodedStreamObject()
    operators = b"BT /F1 12 Tf 20 100 Td (" + text.encode("latin-1", "replace") + b") Tj ET"
    content.set_data(operators)

    # Attach a font dict so the content stream is renderable. pypdf's
    # text-extractor reads stream operators directly so even without a real
    # font registration the (...) Tj string is recovered.
    page[NameObject("/Contents")] = content
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/F1"): DictionaryObject(
                        {
                            NameObject("/Type"): NameObject("/Font"),
                            NameObject("/Subtype"): NameObject("/Type1"),
                            NameObject("/BaseFont"): NameObject("/Helvetica"),
                        }
                    )
                }
            )
        }
    )

    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _make_minimal_docx(text: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_minimal_xlsx(sheets: dict[str, list[list[str]]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    # Drop default sheet, add ours
    default = wb.active
    wb.remove(default)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
